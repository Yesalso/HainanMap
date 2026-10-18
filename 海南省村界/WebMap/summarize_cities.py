# -*- coding: utf-8 -*-
"""summarize_cities.py —— 生成/刷新 Hainan_Code.xlsx 的「市县乡镇汇总」表。

用法：
    python summarize_cities.py

规则：
    1. 以 Sheet1 的 CITY 列分组，统计各县市的 乡 / 镇 / 市 / 区 数量及合计；
       类型按名称后缀判定：名称取 New_Name，为空时回退 TOWN。
    2. 海口市按 HAIKOU_KEEP 列表拆分（与 build_svg_data.py 的 HAOKOU_KEEP 保持一致）：
       - 列表内的单位仍算「海口市」；
       - 海口市内不在列表里的单位，统计上算「琼山市」。
       该拆分**只体现在本汇总表**，不写回 Sheet1，也不改动任何单位的原始所属。

约束：
    - 只读 Sheet1；不新增 Sheet2；汇总表追加/保持在最后。
    - 汇总数值一律写静态值，不写公式。
"""
import collections
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(BASE, "Hainan_Code.xlsx")
SHEET_NAME = "市县乡镇汇总"
TYPES = ("乡", "镇", "市", "区", "其他")

# 统计口径：海口市内以下单位仍算海口市，其余算琼山市（与 build_svg_data.py 一致）
HAIKOU_KEEP = {
    "西秀镇", "长流镇", "海秀镇", "城西镇", "海秀街道", "秀英街道", "人民路街道",
    "海甸街道", "新埠街道", "金宇街道", "中山街道", "滨海街道", "金贸街道",
    "大同街道", "海垦街道", "国兴街道", "海府街道", "博爱街道", "白龙街道",
    "蓝天街道", "和平南街道", "白沙街道", "新海乡",
}
CITY_HAIKOU = "海口市"
CITY_QIONGSHAN = "琼山市"


def classify(name):
    """按末字归类：乡 / 镇 / 市 / 区 / 其他。"""
    s = str(name).strip()
    for k in ("乡", "镇", "市", "区"):
        if s.endswith(k):
            return k
    return "其他"


def stat_city(city, town):
    """统计口径下的所属市县：仅海口市做拆分。"""
    if city == CITY_HAIKOU and str(town).strip() not in HAIKOU_KEEP:
        return CITY_QIONGSHAN
    return city


def read_sheet1(wb):
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    up = [str(c).strip().upper() if c is not None else "" for c in rows[0]]
    ci, ti, ni, cti = (up.index("CODE"), up.index("TOWN"),
                       up.index("NEW_NAME"), up.index("CITY"))
    return rows, (ci, ti, ni, cti)


def aggregate(rows, cols):
    ci, ti, ni, cti = cols
    agg = collections.OrderedDict()
    orig_order = []
    for r in rows[1:]:
        if r[ci] is None:
            continue
        city = stat_city(r[cti], r[ti])
        name = r[ni] if r[ni] else r[ti]
        if r[cti] not in orig_order:
            orig_order.append(r[cti])
        counter = agg.setdefault(city, collections.Counter({k: 0 for k in TYPES}))
        counter[classify(name)] += 1
    return agg, orig_order


def main():
    wb = openpyxl.load_workbook(XLSX_PATH)
    rows, cols = read_sheet1(wb)
    agg, orig_order = aggregate(rows, cols)

    # 行顺序：沿用 Sheet1 中市县首次出现的顺序；海口市后紧随「琼山市」
    order = []
    for city in orig_order:
        if city not in agg:          # 该市县全部被拆分走（理论上不会发生）
            continue
        order.append(city)
        if city == CITY_HAIKOU:
            order.append(CITY_QIONGSHAN)
    if CITY_QIONGSHAN not in agg:
        agg[CITY_QIONGSHAN] = collections.Counter({k: 0 for k in TYPES})

    grand = collections.Counter()
    for c in agg.values():
        grand.update(c)

    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
        for row in ws.iter_rows(min_row=2):
            if row[0].value is not None:
                for cell in row:
                    cell.value = None
    else:
        ws = wb.create_sheet(SHEET_NAME)
        ws.append(["市县"] + list(TYPES) + ["合计"])

    # 表头
    for col, title in enumerate(["市县"] + list(TYPES) + ["合计"], start=1):
        ws.cell(row=1, column=col, value=title)

    r = 2
    for city in order:
        c = agg[city]
        vals = [c[k] for k in TYPES] + [sum(c[k] for k in TYPES)]
        ws.cell(row=r, column=1, value=city)
        for i, v in enumerate(vals, start=2):
            ws.cell(row=r, column=i, value=v)
        r += 1
    # 全省合计
    vals = [grand[k] for k in TYPES] + [sum(grand[k] for k in TYPES)]
    ws.cell(row=r, column=1, value="全省合计")
    for i, v in enumerate(vals, start=2):
        ws.cell(row=r, column=i, value=v)

    # 清理多余的旧行（表行数变少时）
    if ws.max_row > r:
        ws.delete_rows(r + 1, ws.max_row - r)

    # 保证汇总表在最后，且不产生额外 Sheet
    if wb.sheetnames[-1] != SHEET_NAME:
        wb.move_sheet(SHEET_NAME, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(SHEET_NAME))
    wb.save(XLSX_PATH)

    print("汇总表:", SHEET_NAME, "| 行数:", r)
    print("sheets:", wb.sheetnames)
    for city in order:
        c = agg[city]
        print("  %-12s 乡%-4d 镇%-4d 市%-3d 区%-3d 其他%-2d 合计%d"
              % (city, c["乡"], c["镇"], c["市"], c["区"], c["其他"], sum(c.values())))
    print("  %-12s 乡%-4d 镇%-4d 市%-3d 区%-3d 其他%-2d 合计%d"
          % ("全省合计", grand["乡"], grand["镇"], grand["市"], grand["区"],
             grand["其他"], sum(grand.values())))


if __name__ == "__main__":
    main()
