# -*- coding: utf-8 -*-
"""build_election.py —— 从 省长选举.xlsx 抽取选举数据，生成 election_data.js

输出结构（window.ELECTION_DATA）：
    {
      "year": "2010",
      "candidates": [ {"name": "罗保铭", "group": 1}, {"name": "李江华", "group": 2} ],
      "towns": {
         "460105001": {"D": 55203, "E": 72.78, "F": 40177, "G": 39452,
                       "votes": [21523, 17929], "rates": [54.56, 45.44]},
         ...
      }
    }

说明：
  * D = 选举人数（六人普常住人口）  E = 投票率(%)  F = 投票数  G = 有效票
  * votes / rates 的下标与 candidates 顺序一一对应；candidates 顺序 = 表头中
    「得票率」列从左到右的顺序（即 第1组=J、第2组=L、第3组=N、第4组=P）。
  * 只有 **有效票(G) 有值** 的行才写入；其余乡镇在地图侧视为「无数据」。
  * 只读 xlsx，不改动任何表格。
"""
import json
import os

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "省长选举.xlsx")
OUT = os.path.join(BASE, "election_data.js")


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def main():
    if not os.path.exists(XLSX):
        raise SystemExit("找不到文件：%s" % XLSX)

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    up = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

    i_code = up.index("CODE")
    i_pop = up.index("六人普常住人口")
    i_rate = up.index("投票率")
    i_turnout = up.index("投票人数")
    i_valid = up.index("有效票")

    # 候选人 = 表头名为「得票率」的列（按列序 → 组号 1..4），左邻列为该候选人的得票数
    cand_cols = []
    for g, rc in enumerate([i for i, h in enumerate(up) if h == "得票率"], 1):
        cand_cols.append({"name": up[rc - 1] or ("候选人%d" % g), "group": g, "cc": rc - 1, "rc": rc})

    towns = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        code = r[i_code]
        if code is None:
            continue
        if not num(r[i_valid]):
            continue                       # 无有效票 → 无数据，跳过
        votes = [int(r[c["cc"]]) if num(r[c["cc"]]) else None for c in cand_cols]
        rates = [float(r[c["rc"]]) if num(r[c["rc"]]) else None for c in cand_cols]
        if any(v is None for v in votes) or any(v is None for v in rates):
            continue                       # 候选人数据不完整 → 跳过该行
        towns[str(code)] = {
            "D": int(r[i_pop]) if num(r[i_pop]) else None,
            "E": round(float(r[i_rate]), 4) if num(r[i_rate]) else None,
            "F": int(r[i_turnout]) if num(r[i_turnout]) else None,
            "G": int(r[i_valid]),
            "votes": votes,
            "rates": [round(x, 4) for x in rates],
        }

    data = {
        "year": ws.title,
        "source": os.path.basename(XLSX),
        "candidates": [{"name": c["name"], "group": c["group"]} for c in cand_cols],
        "towns": towns,
    }
    js = "window.ELECTION_DATA=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(js)

    print("候选人:", data["candidates"])
    print("有数据的乡镇数:", len(towns))
    print("输出:", OUT, "(%.1f KB)" % (len(js.encode("utf-8")) / 1024))
    for k in list(towns)[:3]:
        print("  样例", k, towns[k])


if __name__ == "__main__":
    main()
