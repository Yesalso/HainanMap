# -*- coding: utf-8 -*-
"""fill_votes.py —— 为「省长选举」表自动填补 投票人数 / 有效票 / 废票，以及各候选人得票数、得票率。

列口径（**按表头名自动定位**，与具体列字母无关；括号内为当前布局）：
    D = 六人普常住人口
    E = 投票率
    F = 投票人数   = D * E / 100
    G = 有效票     = F * k        （k 为有效票占比；同时也是候选人票数之和 = 总票数）
    H = 废票       = F - G

候选人区块（成对出现，形如「票数, 得票率」）：
    I = 罗保铭(票数)   J = 得票率
    K = 李江华(票数)   L = 得票率
    （如继续向右扩展，M/N、O/P… 同理）

    * J / L / N / P … 为「得票率」列，其**左邻列** I / K / M / O … 为对应的「得票数」列。
    * 我们**只写「得票率」**，不写「得票数」——得票数由脚本按 G * 得票率 反推。
    * 若只写了 n 个得票率，说明实际只有 n+1 个候选人；最后一个（余项）取
      **「下一个得票率列」**（不一定在最右侧：只写 J 时，余项就在 L）：
          余项得票率 = 100% - 已列出的得票率之和；
          余项得票数 = G - 已算出的得票数之和。
      即「只有最后一个是前面几个相减」。
    * 若表头的「得票率」列已被**全部写出**（没有空列可放余项），则不再造余项，
      只按 G * 得票率 反推各候选人得票数——即「已写了得票率、但没写得票数」的场景。

随机规则：
    1. F：E 在相乘时**不取显示值**，而是还原成它代表的区间内的随机值
           显示 73.02 -> (73.015, 73.025)；显示 73.1 -> (73.095, 73.105)
           即随机值 ∈ (显示值 - 0.005, 显示值 + 0.005)
    2. G：k 服从**截断正态分布**，取值限定在 [96.9%, 99.2%]
           默认均值 mu = (96.9+99.2)/2 %，标准差 sigma = (99.2-96.9)/6 %（约 ±3σ 覆盖全区间）
    3. 候选人得票数：G × 得票率，其中得票率同样先还原为区间随机值
           (显示值 - 0.005, 显示值 + 0.005)，再四舍五入取整
    4. 余项候选人：得票率 = 100 - Σ(已列得票率)；得票数 = G - Σ(已算得票数)
    5. F / G / 得票数 均四舍五入取整；某行缺 G（无源数据）时整行跳过。

用法：
    python fill_votes.py                     # 原地写回 省长选举.xlsx
    python fill_votes.py --seed 42           # 固定随机种子（可复现）
    python fill_votes.py --dry-run           # 只预览，不写文件
    python fill_votes.py --overwrite         # 已有值也重算覆盖（默认保留已有值）
    python fill_votes.py --half 0.005        # 得票率/投票率 区间半宽
    python fill_votes.py --valid-lo 96.9 --valid-hi 99.2 --sigma 0.3833
    python fill_votes.py --file xx.xlsx --sheet 2010

说明：直接原地补充，不新建文件、不新建文件夹；除上述各列外不动其它列。
"""
import argparse
import math
import os
import random
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_FILE = os.path.join(BASE, "省长选举.xlsx")

COL_POP = ("六人普常住人口", "D")      # 被乘数（人口）
COL_RATE = ("投票率", "E")             # 乘数（百分比）
COL_OUT = ("投票人数", "F")            # F = D*E/100
COL_VALID = ("有效票", "G")            # G = F*k（= 候选人票数之和）
COL_INVALID = ("废票", "H")            # 废票 = F-G（按列名定位，不依赖字母）

RATE_TITLE = "得票率"                  # 候选人「得票率」列的表头名

HALF = 0.005                           # 显示值 -> 区间的半宽（投票率 / 得票率通用）
VALID_LO = 96.9                        # G：有效票占比下限（%）
VALID_HI = 99.2                        # G：有效票占比上限（%）
SIGMA_DIV = 6.0                        # sigma = (hi-lo)/6，即约 ±3σ 覆盖全区间


def round_half_up(x):
    """四舍五入取整（.5 进位，避免 Python 默认的银行家舍入）。"""
    return int(math.floor(x + 0.5))


def rand_rate(shown, half=HALF):
    """把「投票率 / 得票率」显示值还原为区间随机值：strictly (shown-half, shown+half)。"""
    return random.uniform(shown - half, shown + half)


def rand_valid_ratio(lo_pct=VALID_LO, hi_pct=VALID_HI, mu_pct=None, sigma_pct=None):
    """截断正态采样有效票占比，返回小数（如 0.9812）。取值严格落在 [lo, hi] 内。"""
    lo, hi = lo_pct / 100.0, hi_pct / 100.0
    mu = (lo + hi) / 2.0 if mu_pct is None else mu_pct / 100.0
    sigma = (hi - lo) / SIGMA_DIV if sigma_pct is None else sigma_pct / 100.0
    if sigma <= 0:
        return mu
    while True:
        x = random.gauss(mu, sigma)
        if lo <= x <= hi:      # 拒绝采样，保证落在区间内（保留正态形状）
            return x


def pick_col(header, names, fallback_idx):
    """按表头名找列号（0 基）；找不到就用给定的兜底下标。"""
    up = [str(c).strip() if c is not None else "" for c in header]
    for n in names:
        if n in up:
            return up.index(n)
    return fallback_idx


def num(v):
    """是否为可用数值（bool 不算）。"""
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def rate_columns(header):
    """返回所有「得票率」列的 0 基下标，按从左到右排序。"""
    return [i for i, h in enumerate(header)
            if h is not None and str(h).strip() == RATE_TITLE]


def fill_candidates(row, rate_cols, g, overwrite, half, tol=1e-6):
    """填一行的候选人得票数 / 得票率，返回结果字典（无候选人列时返回 None）。

    判定「已写出」的得票率列：从左往右扫描，累计之和首次达到 100 的那一列
    视为**上一轮算出的余项**，它之前的列均为「用户已写出」的候选人；第一个空列
    也视为余项列。余项列之前的各列，按 G * 得票率 反推得票数。

    返回：
        {"rem": 余项得票率列下标或 None, "n": 处理列数, "total": 票数合计,
         "new_cnt": 本次实际写入的得票数单元格数, "detail": [(kind, 票数, 得票率), ...]}
    特例：表头的「得票率」列已被全部写出（没有空列可放余项）时，
          rem=None，只按 G × 得票率 反推各候选人得票数。
    """
    provided = []          # [(rate_idx, count_idx, rate_value)]
    s_rate = 0.0
    rem_rate_idx = None
    for ri in rate_cols:
        v = row[ri].value
        if not num(v):
            rem_rate_idx = ri
            break
        if provided and abs(s_rate + float(v) - 100.0) <= tol:
            rem_rate_idx = ri          # 该列是上一轮写入的余项，不是「已写出」的候选人
            break
        provided.append((ri, ri - 1, float(v)))
        s_rate += float(v)

    # ---- 已写出候选人：得票数 = G * 得票率（得票率先还原为区间随机值）----
    sum_cnt = 0
    new_cnt = 0
    detail = []
    for ri, ci, v in provided:
        cur = row[ci].value
        if overwrite or not num(cur):
            cur = min(round_half_up(g * rand_rate(v, half) / 100.0), g)
            row[ci].value = cur
            new_cnt += 1
        sum_cnt += int(cur)
        detail.append(("prov", row[ci].value, v))

    if rem_rate_idx is None:
        # 没有空列可放余项（表头得票率列已排满）：只反推得票数，不造余项。
        if not detail:
            return None
        return {"rem": None, "n": len(detail), "total": sum_cnt,
                "new_cnt": new_cnt, "detail": detail}

    # ---- 余项候选人：得票率 = 100 - Σ；得票数 = G - Σ ----
    rem_rate = round(100.0 - s_rate, 6)
    rem_cnt = int(g) - sum_cnt
    ri, ci = rem_rate_idx, rem_rate_idx - 1
    if overwrite or not num(row[ri].value):
        row[ri].value = rem_rate
    if overwrite or not num(row[ci].value):
        row[ci].value = rem_cnt
        new_cnt += 1
    detail.append(("rem", row[ci].value, row[ri].value))
    return {"rem": rem_rate_idx, "n": len(provided) + 1, "total": sum_cnt + rem_cnt,
            "new_cnt": new_cnt, "detail": detail}


def main():
    ap = argparse.ArgumentParser(
        description="填补 投票人数(F) / 有效票(G) / 废票(F-G) / 候选人得票数(I,K,…) 与余项得票率(L,…)")
    ap.add_argument("--file", default=DEFAULT_FILE, help="xlsx 路径")
    ap.add_argument("--sheet", default=None, help="工作表名（默认第一个）")
    ap.add_argument("--seed", type=int, default=None, help="随机种子（可选）")
    ap.add_argument("--half", type=float, default=HALF, help="投票率/得票率 区间半宽，默认 0.005")
    ap.add_argument("--valid-lo", type=float, default=VALID_LO, help="有效票占比下限(%%)")
    ap.add_argument("--valid-hi", type=float, default=VALID_HI, help="有效票占比上限(%%)")
    ap.add_argument("--sigma", type=float, default=None,
                    help="有效票占比标准差(%%)，默认 (上限-下限)/6")
    ap.add_argument("--overwrite", action="store_true",
                    help="已有值时也重算覆盖（默认保留已有值）")
    ap.add_argument("--dry-run", action="store_true", help="只预览，不写文件")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        raise SystemExit("找不到文件：%s" % args.file)
    if args.seed is not None:
        random.seed(args.seed)

    wb = openpyxl.load_workbook(args.file)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]

    c_pop = pick_col(header, COL_POP, 3)
    c_rate = pick_col(header, COL_RATE, 4)
    c_out = pick_col(header, COL_OUT, 5)
    c_valid = pick_col(header, COL_VALID, 6)
    c_invalid = pick_col(header, COL_INVALID, 7)
    rcols = rate_columns(header)

    n_f = n_g = n_j = 0
    skipped = 0
    n_cand_rows = 0
    n_cand_filled = 0
    rem_written = {}
    preview = []

    for row in ws.iter_rows(min_row=2):
        pop, rate = row[c_pop].value, row[c_rate].value

        # ---- F：投票人数 = D * E / 100 ----
        if (args.overwrite or not num(row[c_out].value)) \
                and num(pop) and num(rate):
            row[c_out].value = round_half_up(float(pop) * rand_rate(float(rate), args.half) / 100.0)
            n_f += 1

        f = row[c_out].value
        if not num(f):
            skipped += 1
            continue

        # ---- G：有效票 = F * k，k ~ 截断正态[lo, hi] ----
        if args.overwrite or not num(row[c_valid].value):
            k = rand_valid_ratio(args.valid_lo, args.valid_hi, sigma_pct=args.sigma)
            g = round_half_up(float(f) * k)
            g = min(g, int(f))               # 保险：不超过投票人数
            row[c_valid].value = g
            n_g += 1
        g = row[c_valid].value
        if not num(g):
            skipped += 1
            continue

        # ---- H：废票 = F - G ----
        if args.overwrite or not num(row[c_invalid].value):
            row[c_invalid].value = int(f) - int(g)
            n_j += 1

        # ---- 候选人区块：得票数 + 余项得票率 ----
        if rcols:
            res = fill_candidates(row, rcols, int(g), args.overwrite, args.half)
            if res is not None:
                n_cand_rows += 1
                n_cand_filled += res["new_cnt"]
                if res["rem"] is not None:
                    rem_written[res["rem"]] = rem_written.get(res["rem"], 0) + 1
                if len(preview) < 5:
                    preview.append((row[0].value, row[1].value, int(f), int(g), res["detail"]))

    # ---------------- 输出 ----------------
    def colname(i):
        s = ""
        i += 1
        while i:
            i, r = divmod(i - 1, 26)
            s = chr(65 + r) + s
        return s

    print("工作表: %s" % ws.title)
    print("列: 投票人数=%s 有效票=%s 废票=%s | 得票率列: %s"
          % (colname(c_out), colname(c_valid), colname(c_invalid),
             ", ".join(colname(i) for i in rcols) or "无"))
    if rcols:
        rem = sorted(rem_written)
        print("候选人区块: 处理 %d 行 | 得票数列 %s | 余项得票率列 %s | 新填得票数 %d 个"
              % (n_cand_rows,
                 "/".join(colname(i - 1) for i in rcols),
                 "/".join(colname(i) for i in rem) or "无",
                 n_cand_filled))
    print("F 投票人数: 新填 %d | G 有效票: 新填 %d | 废票: 新填 %d | 跳过 %d 行"
          % (n_f, n_g, n_j, skipped))
    print("有效票占比: 截断正态 [%.2f%%, %.2f%%], mu=%.4f%%, sigma=%s"
          % (args.valid_lo, args.valid_hi, (args.valid_lo + args.valid_hi) / 2,
             ("%.4f%%" % args.sigma) if args.sigma is not None
             else "(%.4f%% of (hi-lo)/6)" % ((args.valid_hi - args.valid_lo) / SIGMA_DIV)))
    if preview:
        print("样例（CODE, 乡镇, F, G, 各候选人[票数,得票率]）:")
        for code, town, f, g, detail in preview:
            seg = " ".join("[%s, %s]" % (a, b) for _, a, b in detail)
            print("    %s %s F=%d G=%d -> %s" % (code, town, f, g, seg))

    if args.dry_run:
        print("--dry-run：未写入文件")
        return
    wb.save(args.file)
    print("已写回：%s" % args.file)


if __name__ == "__main__":
    main()
